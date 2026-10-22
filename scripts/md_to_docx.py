# -*- coding: utf-8 -*-
"""将课程报告 Markdown 转为可提交的 docx。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm


def set_run_font(run, name: str = '宋体', size: int = 12, bold: bool = False):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = name
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading_cn(doc: Document, text: str, level: int):
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, '黑体', 18, True)
    elif level == 1:
        run = p.add_run(text)
        set_run_font(run, '黑体', 14, True)
    else:
        run = p.add_run(text)
        set_run_font(run, '黑体', 12, True)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_run_font(run, '宋体', 12)
    return p


def main():
    root = Path(__file__).resolve().parents[1]
    md_path = root / 'docs' / '课程报告.md'
    out_path = root / 'docs' / '课程报告-出行天气助手.docx'
    lines = md_path.read_text(encoding='utf-8').splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith('# '):
            add_heading_cn(doc, line[2:].strip(), 0)
        elif line.startswith('## '):
            add_heading_cn(doc, line[3:].strip(), 1)
        elif line.startswith('### '):
            add_heading_cn(doc, line[4:].strip(), 2)
        elif line.startswith('|') or line.startswith('```') or line.startswith('---'):
            # 表格/代码块简化为正文
            if line.startswith('```'):
                continue
            p = doc.add_paragraph()
            run = p.add_run(line)
            set_run_font(run, 'Consolas', 10)
        elif line.startswith('> '):
            p = doc.add_paragraph()
            run = p.add_run(line[2:])
            set_run_font(run, '楷体', 12)
        elif line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(line[2:])
            set_run_font(run, '宋体', 12)
        else:
            add_body(doc, line.lstrip('- ').lstrip('> ').strip())

    doc.save(out_path)
    print('wrote', out_path)


if __name__ == '__main__':
    main()
